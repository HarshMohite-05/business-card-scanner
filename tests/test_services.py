import concurrent.futures
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from PIL import Image, ImageDraw, ImageFont
from openpyxl import load_workbook
from services.database_service import ContactStore
from services.contact_parser import parse_contact
from services.image_service import validate_image, preprocess
from services.ocr_service import read_card

class StoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.store = ContactStore(self.root / 'contacts.db', self.root / 'Contacts.xlsx')
        self.store.initialize()

    def tearDown(self):
        self.temp.cleanup()

    def workbook(self):
        return load_workbook(self.store.excel_path)

    def test_empty_startup(self):
        self.assertEqual(self.store.categories(), [])
        self.assertEqual(self.store.contacts(), [])
        workbook = self.workbook()
        sheet = workbook['Contacts']
        self.assertEqual(sheet.max_column, 21)
        self.assertEqual(sheet['D1'].value, 'Added By')
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertTrue(sheet['A1'].font.bold)
        self.assertEqual(sheet.max_row, 1)
        workbook.close()

    def test_lifecycle_and_restart(self):
        first = self.store.save_contact({'name': 'Rajesh Shah', 'email': 'raj@example.com', 'city': 'Mumbai', 'remark': 'Met at exhibition', 'added_by': 'Deepak'}, ' Industrial   Client ')
        second = self.store.save_contact({'company_name': 'Second Company'}, 'industrial client')
        self.assertEqual(len(self.store.categories()), 1)
        self.assertEqual(self.store.categories()[0]['name'], 'Industrial Client')
        self.assertEqual(len(self.store.contacts('MUMBAI', self.store.categories()[0]['id'])), 1)
        self.assertEqual(len(self.store.contacts('exhibition')), 1)
        self.assertEqual(len(self.store.contacts('deepak')), 1)
        contact = self.store.get_contact(first)
        self.store.save_contact({**contact, 'name': 'Corrected Name'}, 'Consultant', contact_id=first)
        self.assertEqual(len(self.store.contacts()), 2)
        workbook = self.workbook()
        self.assertEqual(workbook['Contacts']['D2'].value, 'Deepak')
        self.assertEqual(workbook['Contacts']['E2'].value, 'Corrected Name')
        self.assertEqual(workbook['Contacts']['S2'].value, 'Consultant')
        workbook.close()
        self.store.rename_category(self.store.get_contact(first)['category_id'], 'Adviser')
        self.assertEqual(self.store.get_contact(first)['category'], 'Adviser')
        self.store.delete_contact(first)
        self.assertEqual(self.store.contacts()[0]['id'], second)
        self.assertEqual(self.store.contacts()[0]['serial_number'], 1)
        workbook = self.workbook()
        self.assertEqual(workbook['Contacts'].max_row, 2)
        workbook.close()
        reopened = ContactStore(self.store.db_path, self.store.excel_path)
        reopened.initialize()
        self.assertEqual(reopened.contacts()[0]['id'], second)
        self.assertEqual(len(reopened.categories()), 2)
        self.assertTrue(reopened.export_excel().startswith(b'PK'))

    def test_validation_and_categories(self):
        with self.assertRaises(ValueError):
            self.store.save_contact({}, 'Supplier')
        with self.assertRaises(ValueError):
            self.store.save_contact({'name': 'Test'}, '')
        with self.assertRaises(ValueError):
            self.store.save_contact({'name': 'Test', 'added_by': 'Someone Else'}, 'Supplier')
        self.assertEqual(self.store.categories(), [])
        contact_id = self.store.save_contact({'name': 'Test'}, 'Supplier')
        cat_id = self.store.get_contact(contact_id)['category_id']
        with self.assertRaises(ValueError):
            self.store.delete_category(cat_id)
        self.store.delete_contact(contact_id)
        self.store.delete_category(cat_id)
        self.assertEqual(self.store.categories(), [])

    def test_failed_export_rolls_back(self):
        self.store.save_contact({'name': 'Existing'}, 'Client')
        with patch('services.database_service.replace_workbook', side_effect=PermissionError('locked')):
            with self.assertRaises(PermissionError):
                self.store.save_contact({'name': 'New'}, 'Supplier')
        self.assertEqual(len(self.store.contacts()), 1)
        self.assertEqual(len(self.store.categories()), 1)

    def test_concurrency_and_text_safety(self):
        with concurrent.futures.ThreadPoolExecutor(max_workers=4) as pool:
            list(pool.map(lambda i: self.store.save_contact({'name': f'Person {i}', 'remark': '=1+1'}, 'Client'), range(8)))
        self.assertEqual(len(self.store.contacts()), 8)
        workbook = self.workbook()
        self.assertEqual(workbook['Contacts'].max_row, 9)
        self.assertEqual(workbook['Contacts']['T2'].data_type, 's')
        workbook.close()
        contact = self.store.contacts()[0]
        self.store.save_contact({**contact, 'name': 'Updated'}, 'Client', contact_id=contact['id'])
        with self.assertRaises(ValueError):
            self.store.save_contact(contact, 'Client', contact_id=contact['id'], expected_updated_at=contact['updated_at'])

    def test_legacy_schema_migration_and_two_sides(self):
        import sqlite3
        db = sqlite3.connect(self.store.db_path)
        db.execute('ALTER TABLE contacts DROP COLUMN business_card_back_image')
        db.commit()
        db.close()
        self.store.initialize()
        with patch('services.image_service.IMAGE_ROOT', (self.root / 'images').resolve()):
            from services.image_service import image_path
            front, back = BytesIO(), BytesIO()
            Image.new('RGB', (800, 400), 'white').save(front, format='PNG')
            Image.new('RGB', (800, 400), 'blue').save(back, format='PNG')
            cid = self.store.save_contact({'name': 'Two Sides'}, 'Client', image_data=front.getvalue(), back_image_data=back.getvalue())
            contact = self.store.get_contact(cid)
            paths = [image_path(contact[key]) for key in ('business_card_image', 'business_card_back_image')]
            self.assertEqual(paths[0].read_bytes(), front.getvalue())
            self.assertEqual(paths[1].read_bytes(), back.getvalue())
            self.store.save_contact({**contact, 'name': 'Edited'}, 'Client', contact_id=cid)
            self.assertEqual(self.store.get_contact(cid)['business_card_back_image'], contact['business_card_back_image'])
            workbook = self.workbook()
            self.assertIn(contact['business_card_back_image'], workbook['Contacts']['U2'].value)
            self.assertIn(contact['business_card_image'], workbook['Contacts']['U2'].value)
            workbook.close()
            self.store.delete_contact(cid)
            self.assertTrue(all(not path.exists() for path in paths))

class OCRTests(unittest.TestCase):
    def test_parser(self):
        text = 'Rajesh Shah\nPurchase Manager\nABC Chemicals Pvt Ltd\n+91 98765 43210\nMobile: 9988776655\nTel: 079-23456789\nrajesh@abc.com\nsales@abc.com\nwww.abc.com\nVatva GIDC, Ahmedabad, Gujarat 382445\nIndia'
        contact = parse_contact(text)
        self.assertEqual(contact['name'], 'Rajesh Shah')
        self.assertEqual(contact['email'], 'rajesh@abc.com')
        self.assertEqual(contact['alternate_email'], 'sales@abc.com')
        self.assertEqual(contact['website'], 'www.abc.com')
        self.assertEqual(contact['city'], 'Ahmedabad')
        self.assertEqual(contact['pin_code'], '382445')
        self.assertEqual(contact['telephone_number'], '079-23456789')
        self.assertEqual(contact['alternate_mobile'], '9988776655')
        partial = parse_contact('sales@company.co.in')
        self.assertEqual(partial['email'], 'sales@company.co.in')
        self.assertEqual(partial['website'], '')
        self.assertEqual(partial['country'], '')
        self.assertTrue(all(not v for v in parse_contact('').values()))

    def test_multiline_address_and_spaced_email(self):
        result = parse_contact('MAHESH KAVADKAR\nENGINEER\nservice. mum@example.com\n601 B, Neelyog Square, R.B Mehta\nMarg, Ghatkopar East, Mumbai,\nMaharashtra - 400077, India')
        self.assertEqual(result['email'], 'service.mum@example.com')
        self.assertIn('601 B', result['full_address'])
        self.assertIn('Marg', result['full_address'])
        self.assertEqual(result['city'], 'Mumbai')
        self.assertEqual(result['pin_code'], '400077')

    def test_image_validation(self):
        with self.assertRaises(ValueError):
            validate_image(b'not an image', 'bad.png')
        image = Image.new('RGB', (800, 400), 'white')
        data = BytesIO()
        image.save(data, format='PNG')
        self.assertEqual(validate_image(data.getvalue(), 'test.png').size, (800, 400))
        with self.assertRaises(ValueError):
            validate_image(data.getvalue(), 'test.exe')
        self.assertEqual(preprocess(data.getvalue()).mode, 'L')

    def test_merge_complementary_details(self):
        from services.contact_parser import merge_contacts
        front = parse_contact('Alice Example\nDirector\nEXAMPLE WATER\n+91 9876543210\nalice@example.com')
        back = parse_contact('Example Water Ltd\n98765 43210\nsales@example.com\nwww.example.com\n1 Main Road, Mumbai, Maharashtra 400001\nIndia')
        merged, alternatives = merge_contacts([front, back])
        self.assertEqual(merged['name'], 'Alice Example')
        self.assertEqual(merged['email'], 'alice@example.com')
        self.assertEqual(merged['alternate_email'], 'sales@example.com')
        self.assertEqual(merged['website'], 'www.example.com')
        self.assertEqual(merged['alternate_mobile'], '')
        self.assertEqual(merged['city'], 'Mumbai')
        self.assertIn('company_name', alternatives)

    def test_noisy_card_readings(self):
        from services.contact_parser import merge_contacts, company_from_domain
        icons = parse_contact('Sunil Pagare\nMarketing\n| + 9186898 17578\n\' +91 90046 09111\n© Office no-1, Shivaji nagar, Thane West -400602\nPAR\nTNER')
        self.assertEqual(icons['mobile_number'], '+91 86898 17578')
        self.assertEqual(icons['full_address'], 'Office no-1, Shivaji nagar, Thane West -400602')
        self.assertEqual((icons['city'], icons['state']), ('Thane', 'Maharashtra'))
        self.assertEqual(parse_contact('Vatsal Doshi\nPAR\nTNER')['designation'], 'PARTNER')
        # Misread copies across passes collapse; distinct numbers printed together stay.
        passes = [{'email': 'vatsal@aquatatva.com', 'mobile_number': '+91 99337 74468'},
                  {'email': 'vatszl@aquatatva.com', 'mobile_number': '+91 99337 74468', 'telephone_number': '+9129337 74468'},
                  {'email': 'vatsal@aquatatva.com', 'mobile_number': '9876543210', 'alternate_mobile': '9876543211'}]
        merged, _ = merge_contacts(passes, weights=[70, 60, 40])
        self.assertEqual(merged['email'], 'vatsal@aquatatva.com')
        self.assertEqual(merged['alternate_email'], '')
        self.assertEqual((merged['mobile_number'], merged['alternate_mobile']), ('+91 99337 74468', '9876543210'))
        self.assertEqual(merged['telephone_number'], '9876543211')
        # A logo-only side must not override the details side when weighted.
        merged, _ = merge_contacts([{'name': 'UATAT VA'}, {'name': 'Vatsal Doshi'}], weights=[42, 113])
        self.assertEqual(merged['name'], 'Vatsal Doshi')
        merged = {'company_name': 'ATATVA', 'email': 'vatsal@aquatatva.com', 'website': ''}
        company_from_domain(merged, {})
        self.assertEqual(merged['company_name'], 'AQUATATVA')

    def test_ocr_pass_failure_keeps_other_results(self):
        from services.ocr_service import TesseractEngine, OCRError
        engine = TesseractEngine()
        with patch.object(engine, '_read', side_effect=[
            {'raw_text': 'Alice Example', 'confidence': 90},
            OCRError('timeout'),
            {'raw_text': 'sales@example.com', 'confidence': 90},
            {'raw_text': '', 'confidence': 0},
            {'raw_text': '', 'confidence': 0},
            {'raw_text': '', 'confidence': 0}]):
            output = engine.read(Image.new('L', (800, 400), 'white'))
        self.assertEqual(output['fields']['name'], 'Alice Example')
        self.assertEqual(output['fields']['email'], 'sales@example.com')

    def test_crop_copy_preserves_original(self):
        from services.image_service import processing_copy
        image = Image.new('RGB', (1000, 600), 'white')
        output = BytesIO()
        image.save(output, format='PNG')
        original = output.getvalue()
        cropped = processing_copy(original, (10, 20, 90, 80), 90)
        self.assertEqual(validate_image(cropped).size, (360, 800))
        self.assertEqual(validate_image(original).size, (1000, 600))
        with self.assertRaises(ValueError):
            processing_copy(original, (90, 0, 10, 100))

    def test_original_image_preserved_and_cleanup(self):
        from services.image_service import image_path
        image = Image.new('RGB', (800, 400), 'white')
        output = BytesIO()
        image.save(output, format='PNG')
        data = output.getvalue()
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            store = ContactStore(root / 'contacts.db', root / 'Contacts.xlsx')
            store.initialize()
            with patch('services.image_service.IMAGE_ROOT', root / 'images'):
                cid = store.save_contact({'name': 'Photo Contact'}, 'Client', image_data=data)
                contact = store.get_contact(cid)
                path = image_path(contact['business_card_image'])
                self.assertEqual(path.read_bytes(), data)
                store.delete_contact(cid)
                self.assertFalse(path.exists())

    def test_real_ocr(self):
        image = Image.new('RGB', (1400, 700), 'white')
        draw = ImageDraw.Draw(image)
        font = ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf', 44) if Path('/System/Library/Fonts/Supplemental/Arial.ttf').exists() else ImageFont.truetype('DejaVuSans.ttf', 44)
        lines = ['Rajesh Shah', 'Purchase Manager', 'ABC Chemicals Pvt Ltd', '+91 98765 43210', 'rajesh@abc.com', 'www.abc.com', 'Ahmedabad, Gujarat 382445']
        for i, line in enumerate(lines):
            draw.text((70, 45 + i * 85), line, fill='black', font=font)
        data = BytesIO()
        image.save(data, format='PNG')
        output = read_card(data.getvalue())
        self.assertIn('rajesh@abc.com', output['raw_text'])
        contact = parse_contact(output['raw_text'])
        self.assertEqual(contact['name'], 'Rajesh Shah')
        self.assertEqual(contact['company_name'], 'ABC Chemicals Pvt Ltd')
        blank = BytesIO()
        Image.new('RGB', (800, 400), 'white').save(blank, format='PNG')
        self.assertEqual(read_card(blank.getvalue())['raw_text'], '')

if __name__ == '__main__':
    unittest.main()
