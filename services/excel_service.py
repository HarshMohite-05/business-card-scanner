from datetime import date, time
from io import BytesIO
import os
from pathlib import Path
import tempfile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from utils.constants import EXCEL_COLUMNS, EXCEL_PATH

WIDTHS = {'serial_number': 9, 'date_added': 15, 'time_added': 15, 'added_by': 14, 'name': 25, 'designation': 28, 'company_name': 32,
          'mobile_number': 24, 'alternate_mobile': 24, 'telephone_number': 24, 'email': 32, 'alternate_email': 32, 'website': 32,
          'full_address': 48, 'city': 22, 'state': 22, 'pin_code': 13, 'country': 22, 'category': 28, 'remark': 48, 'business_card_image': 55}
WRAPPED = {i for i, (field, _) in enumerate(EXCEL_COLUMNS, 1) if field in ('full_address', 'remark')}


def build_workbook(contacts):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Contacts'
    sheet.append([label for _, label in EXCEL_COLUMNS])
    for serial, contact in enumerate(contacts, 1):
        values = []
        for field, _ in EXCEL_COLUMNS:
            value = serial if field == 'serial_number' else contact.get(field, '')
            if field == 'business_card_image':
                value = '; '.join(v for v in (contact.get('business_card_image'), contact.get('business_card_back_image')) if v)
            if field == 'date_added' and value:
                value = date.fromisoformat(value)
            elif field == 'time_added' and value:
                value = time.fromisoformat(value)
            elif isinstance(value, str):
                value = ILLEGAL_CHARACTERS_RE.sub('', value)
            values.append(value)
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            # Untrusted contact values must remain text, including strings beginning with '='.
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=cell.column in WRAPPED)
        sheet.cell(sheet.max_row, 2).number_format = 'dd-mm-yyyy'
        sheet.cell(sheet.max_row, 3).number_format = 'hh:mm:ss'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='244A73')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    for i, (field, _) in enumerate(EXCEL_COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(i)].width = WIDTHS.get(field, 20)
    sheet.row_dimensions[1].height = 34
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(EXCEL_COLUMNS))}{max(1, sheet.max_row)}'
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def replace_workbook(data, path=EXCEL_PATH):
    """Validate the temporary file before an atomic replacement on the same filesystem."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix='.contacts_', suffix='.xlsx', dir=path.parent)
    try:
        with os.fdopen(fd, 'wb') as file:
            file.write(data)
            file.flush()
            os.fsync(file.fileno())
        check = load_workbook(temporary, read_only=True)
        try:
            if check.sheetnames != ['Contacts'] or check['Contacts'].max_column != len(EXCEL_COLUMNS):
                raise ValueError('Excel workbook validation failed.')
        finally:
            check.close()
        os.replace(temporary, path)
    finally:
        Path(temporary).unlink(missing_ok=True)


def regenerate_excel(store):
    return store.export_excel()
